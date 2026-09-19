"""
Exporting the remittance report.

Both exports read `_remittance_report_data`, the same helper the page renders
from, so a downloaded figure cannot drift from the one on screen. These check
the branch/date filters carry through and that the breakdowns are included.
"""
from datetime import timedelta
from io import BytesIO

import openpyxl
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from autodash_App.models import (
    Branch, CustomUser, DailyRemittance, RemittancePayment, RemittanceSetup,
    Worker, WorkerCategory,
)


class RemittanceExportTest(TestCase):

    def setUp(self):
        self.ridge = Branch.objects.create(
            name="Ridge", location="Accra", phone_number="0240000000"
        )
        self.teshie = Branch.objects.create(
            name="Teshie", location="Accra", phone_number="0240000001"
        )
        RemittanceSetup.objects.create(branch=self.ridge, target_amount=1200.0)
        RemittanceSetup.objects.create(branch=self.teshie, target_amount=800.0)
        self.today = timezone.localdate()

        self.ridge_row = DailyRemittance.objects.create(
            branch=self.ridge, date=self.today - timedelta(days=1),
            target_amount=1200.0, net_sales=1500.0, is_finalized=True,
        )
        self.ridge_row.recalc()
        RemittancePayment.objects.create(
            remittance=self.ridge_row, source="cash", amount=1200.0
        )
        self.teshie_row = DailyRemittance.objects.create(
            branch=self.teshie, date=self.today - timedelta(days=1),
            target_amount=800.0, net_sales=600.0, is_finalized=True,
        )
        self.teshie_row.recalc()
        RemittancePayment.objects.create(
            remittance=self.teshie_row, source="momo", amount=500.0
        )

        self.staff = CustomUser.objects.create_user(
            username="0248888888", password="x", role="worker",
            is_staff=True, is_superuser=True, approved=True,
        )
        self.client.force_login(self.staff)

    def _sheet(self, **params):
        response = self.client.get(reverse("remittance_report_excel"), params)
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        return response, wb.active

    def _cells(self, ws):
        return [
            [c for c in row]
            for row in ws.iter_rows(values_only=True)
        ]

    # ---- excel ----------------------------------------------------------
    def test_excel_downloads_as_a_spreadsheet(self):
        response, _ = self._sheet()
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_excel_covers_every_branch_by_default(self):
        _, ws = self._sheet()
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertIn("All branches", flat)
        self.assertIn("Ridge", flat)
        self.assertIn("Teshie", flat)

    def test_excel_carries_the_branch_filter(self):
        _, ws = self._sheet(branch_id=self.ridge.id)
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertIn("Ridge", flat)
        self.assertNotIn("Teshie", flat)

    def test_excel_includes_the_branch_breakdown_when_unfiltered(self):
        _, ws = self._sheet()
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertIn("Breakdown by branch", flat)

    def test_excel_drops_the_branch_breakdown_when_filtered(self):
        _, ws = self._sheet(branch_id=self.ridge.id)
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertNotIn("Breakdown by branch", flat)

    def test_excel_includes_the_source_breakdown(self):
        _, ws = self._sheet()
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertIn("Breakdown by source", flat)

    def test_excel_totals_match_the_page(self):
        page = self.client.get(reverse("remittance_report"))
        _, ws = self._sheet()
        rows = self._cells(ws)
        totals_row = next(r for r in rows if r and r[0] == "Totals")
        # Unfiltered layout: Totals, (branch), net, target, '', due, remitted, ...
        self.assertAlmostEqual(totals_row[2], page.context["totals"]["net_sales"])
        self.assertAlmostEqual(totals_row[6], page.context["totals"]["remitted"])

    def test_excel_names_a_carried_day_as_carried(self):
        self.teshie_row.propagate_forward()
        DailyRemittance.objects.create(
            branch=self.teshie, date=self.today, target_amount=800.0,
            brought_forward=self.teshie_row.carry_to_next_day(),
        ).recalc()
        self.teshie_row.propagate_forward()
        _, ws = self._sheet(branch_id=self.teshie.id)
        flat = " ".join(str(c) for row in self._cells(ws) for c in row if c)
        self.assertIn("Carried", flat)

    def test_excel_handles_an_empty_range(self):
        far = (self.today - timedelta(days=400)).strftime("%Y-%m-%d")
        response, ws = self._sheet(start_date=far, end_date=far)
        self.assertEqual(response.status_code, 200)

    # ---- pdf ------------------------------------------------------------
    def _pdf(self, **params):
        response = self.client.get(reverse("remittance_report_pdf"), params)
        self.assertEqual(response.status_code, 200)
        return response

    def test_pdf_downloads_as_a_pdf(self):
        response = self._pdf()
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment;", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_pdf_filename_names_the_scope_and_range(self):
        self.assertIn("all_branches", self._pdf()["Content-Disposition"])
        self.assertIn(
            "ridge", self._pdf(branch_id=self.ridge.id)["Content-Disposition"]
        )

    def test_pdf_handles_an_empty_range(self):
        far = (self.today - timedelta(days=400)).strftime("%Y-%m-%d")
        response = self._pdf(start_date=far, end_date=far)
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_pdf_is_not_empty(self):
        self.assertGreater(len(self._pdf().content), 1000)


class RemittanceExportAccessTest(TestCase):
    """The exports are gated exactly like the report they come from."""

    def setUp(self):
        self.ridge = Branch.objects.create(
            name="Ridge", location="Accra", phone_number="0240000000"
        )
        self.teshie = Branch.objects.create(
            name="Teshie", location="Accra", phone_number="0240000001"
        )
        DailyRemittance.objects.create(
            branch=self.teshie, date=timezone.localdate(), target_amount=800.0
        ).recalc()

    def test_anonymous_is_turned_away(self):
        for name in ("remittance_report_excel", "remittance_report_pdf"):
            response = self.client.get(reverse(name))
            self.assertIn(response.status_code, (302, 403), msg=name)

    def test_a_plain_worker_is_turned_away(self):
        category = WorkerCategory.objects.create(name="Washer", service_provider=True)
        user = CustomUser.objects.create_user(
            username="0247777777", password="x", role="worker", approved=True
        )
        Worker.objects.create(user=user, branch=self.ridge, worker_category=category)
        self.client.force_login(user)
        for name in ("remittance_report_excel", "remittance_report_pdf"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403, msg=name)

    def test_a_branch_admin_is_pinned_to_their_own_branch(self):
        category = WorkerCategory.objects.create(name="Lead", service_provider=False)
        user = CustomUser.objects.create_user(
            username="0246666666", password="x", role="worker", approved=True
        )
        Worker.objects.create(
            user=user, branch=self.ridge, worker_category=category,
            is_branch_admin=True,
        )
        self.client.force_login(user)
        response = self.client.get(
            reverse("remittance_report_excel"), {"branch_id": self.teshie.id}
        )
        self.assertEqual(response.status_code, 200)
        wb = openpyxl.load_workbook(BytesIO(response.content))
        flat = " ".join(
            str(c) for row in wb.active.iter_rows(values_only=True) for c in row if c
        )
        self.assertIn("Ridge", flat)
        self.assertNotIn("Teshie", flat)
